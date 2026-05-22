"""Backend test for HydroManager Excel export endpoints."""
import os
import io
import re
import sys
import json
import requests
from openpyxl import load_workbook
from datetime import date, timedelta

BASE = "https://hydro-check-log.preview.emergentagent.com/api"
ADMIN_EMAIL = "akhilharisai@gmail.com"
ADMIN_PASSWORD = "Admin@123"

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

results = []  # list of (name, passed, message)


def record(name, passed, msg=""):
    results.append((name, passed, msg))
    flag = "PASS" if passed else "FAIL"
    print(f"[{flag}] {name} :: {msg}")


def login_admin():
    r = requests.post(f"{BASE}/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD, "device": "backend-test"},
                      timeout=30)
    r.raise_for_status()
    data = r.json()
    otp = data.get("dev_otp")
    assert otp, f"No dev_otp returned: {data}"
    r2 = requests.post(f"{BASE}/auth/verify-otp",
                       json={"email": ADMIN_EMAIL, "otp": otp},
                       timeout=30)
    r2.raise_for_status()
    return r2.json()["token"]


def seed_data(token, today):
    h = {"Authorization": f"Bearer {token}"}
    # Tank reading
    tank_payload = {
        "check_date": today, "check_time": "08:30", "session": "morning", "tank_id": 1,
        "ph_actual": 6.1, "ph_target_min": 5.8, "ph_target_max": 6.4,
        "ec_actual": 1.5, "ec_target_min": 1.2, "ec_target_max": 1.8,
        "temp_actual": 22.5, "temp_target_min": 18.0, "temp_target_max": 25.0,
        "notes": "Backend export test - tank 1"
    }
    r = requests.post(f"{BASE}/tanks/reading", json=tank_payload, headers=h, timeout=20)
    record("seed: POST /tanks/reading", r.status_code == 200, f"status={r.status_code} body={r.text[:200]}")

    env_payload = {
        "check_date": today, "check_time": "09:00", "session": "morning",
        "temperature": 24.5, "humidity": 65.0, "notes": "Backend export test env"
    }
    r = requests.post(f"{BASE}/environment/reading", json=env_payload, headers=h, timeout=20)
    record("seed: POST /environment/reading", r.status_code == 200, f"status={r.status_code} body={r.text[:200]}")

    field_payload = {
        "check_date": today,
        "seedling_watered": True, "seedling_ph": 6.0, "seedling_ec": 1.0,
        "pest_check_done": True, "pest_notes": "No pests observed",
        "leaf_cleaning_done": False, "leaf_notes": "",
        "notes": "Backend export test field"
    }
    r = requests.post(f"{BASE}/field/tasks", json=field_payload, headers=h, timeout=20)
    record("seed: POST /field/tasks", r.status_code == 200, f"status={r.status_code} body={r.text[:200]}")

    weekly_payload = {
        "check_date": today,
        "meter_calibration_done": True,
        "nutrition_quantity_ok": True,
        "nutrition_notes": "Weekly stock fine",
        "tank_filters_cleaned": True,
        "notes": "Backend export test weekly"
    }
    r = requests.post(f"{BASE}/checks/weekly", json=weekly_payload, headers=h, timeout=20)
    record("seed: POST /checks/weekly", r.status_code == 200, f"status={r.status_code} body={r.text[:200]}")

    monthly_payload = {
        "check_date": today,
        "tanks_cleaned": True,
        "salt_formation_ok": True,
        "salt_notes": "Looks clear",
        "solution_a_qty": 12.5,
        "solution_b_qty": 10.0,
        "solution_c_qty": 8.5,
        "solutions_ordered": False,
        "seeds_qty_ok": True,
        "seeds_ordered": True,
        "notes": "Backend export test monthly"
    }
    r = requests.post(f"{BASE}/checks/monthly", json=monthly_payload, headers=h, timeout=20)
    record("seed: POST /checks/monthly", r.status_code == 200, f"status={r.status_code} body={r.text[:200]}")


def test_export(token, path_segment, friendly, start, end,
                expected_filename_prefix,
                expected_sheets,  # list of expected sheet titles to be present
                expected_headers,  # dict sheet -> list of expected header substrings (must include all)
                row_marker=None):
    """Test one export endpoint comprehensively."""
    h = {"Authorization": f"Bearer {token}"}
    url = f"{BASE}/export/{path_segment}"

    # 1) Unauthenticated
    r0 = requests.get(url, timeout=30)
    record(f"{friendly}: unauth returns 401/403",
           r0.status_code in (401, 403),
           f"status={r0.status_code}")

    # 2) Authenticated WITH date range
    r = requests.get(url, params={"start": start, "end": end}, headers=h, timeout=60)
    record(f"{friendly}: GET with date range -> 200",
           r.status_code == 200,
           f"status={r.status_code} body={r.text[:200] if r.status_code != 200 else 'ok'}")
    if r.status_code != 200:
        return

    # 3) Content-Type
    ct = r.headers.get("content-type", "")
    record(f"{friendly}: Content-Type is xlsx",
           XLSX_MIME in ct,
           f"got '{ct}'")

    # 4) Content-Disposition filename
    cd = r.headers.get("content-disposition", "")
    expected_fn = f'{expected_filename_prefix}_{start}_to_{end}.xlsx'
    record(f"{friendly}: Content-Disposition filename matches",
           expected_fn in cd,
           f"expected '{expected_fn}' in '{cd}'")

    # 5) Valid xlsx + sheets/headers
    try:
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
        sheet_names = wb.sheetnames
        for s in expected_sheets:
            record(f"{friendly}: contains sheet '{s}'",
                   s in sheet_names,
                   f"sheets={sheet_names}")
        for s, header_list in expected_headers.items():
            if s not in sheet_names:
                continue
            ws = wb[s]
            row1 = [c.value for c in ws[1]]
            missing = [h for h in header_list if h not in row1]
            record(f"{friendly}: sheet '{s}' has expected headers",
                   len(missing) == 0,
                   f"row1={row1}  missing={missing}")
            # Header count expectation: store actual length
            print(f"    INFO {friendly}/{s} header count = {len(row1)}")
        # row marker check (data row exists)
        if row_marker is not None:
            sheet, marker = row_marker
            ws = wb[sheet]
            found = False
            for row in ws.iter_rows(min_row=2, values_only=True):
                for cell in row:
                    if cell is not None and marker in str(cell):
                        found = True
                        break
                if found:
                    break
            record(f"{friendly}: seeded row appears in '{sheet}'",
                   found,
                   f"marker='{marker}'")
    except Exception as e:
        record(f"{friendly}: xlsx parse error", False, str(e))

    # 6) No params -> still valid xlsx
    r2 = requests.get(url, headers=h, timeout=60)
    record(f"{friendly}: GET without params -> 200",
           r2.status_code == 200,
           f"status={r2.status_code} body={r2.text[:200] if r2.status_code != 200 else 'ok'}")
    if r2.status_code == 200:
        try:
            wb2 = load_workbook(io.BytesIO(r2.content), data_only=True)
            record(f"{friendly}: no-params response is valid xlsx", True, f"sheets={wb2.sheetnames}")
        except Exception as e:
            record(f"{friendly}: no-params response is valid xlsx", False, str(e))


def main():
    print(f"Backend base: {BASE}")
    try:
        token = login_admin()
        record("auth: login + verify-otp", True, "obtained token")
    except Exception as e:
        record("auth: login + verify-otp", False, str(e))
        print_summary()
        sys.exit(1)

    today = date.today().isoformat()
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    tomorrow = (date.today() + timedelta(days=1)).isoformat()

    seed_data(token, today)

    # Tanks
    test_export(
        token, "tanks", "/api/export/tanks",
        start=yesterday, end=tomorrow,
        expected_filename_prefix="Tanks",
        expected_sheets=["Tank Readings"],
        expected_headers={"Tank Readings": [
            "Date", "Time", "Session", "Tank", "pH",
            "pH Target Min", "pH Target Max",
            "EC", "EC Target Min", "EC Target Max",
            "Temp (°C)", "Temp Target Min", "Temp Target Max",
            "Notes", "Recorded By", "Created At"
        ]},
        row_marker=("Tank Readings", "Backend export test - tank 1"),
    )

    # Environment
    test_export(
        token, "environment", "/api/export/environment",
        start=yesterday, end=tomorrow,
        expected_filename_prefix="Environment",
        expected_sheets=["Environment", "Daily Averages"],
        expected_headers={
            "Environment": ["Date", "Time", "Session", "Temperature (°C)", "Humidity (%)", "Notes", "Recorded By", "Created At"],
            "Daily Averages": ["Date", "Avg Temperature (°C)", "Avg Humidity (%)", "Readings"],
        },
        row_marker=("Environment", "Backend export test env"),
    )

    # Field
    test_export(
        token, "field", "/api/export/field",
        start=yesterday, end=tomorrow,
        expected_filename_prefix="Field",
        expected_sheets=["Field Tasks"],
        expected_headers={"Field Tasks": [
            "Date", "Seedling Watered", "Seedling pH", "Seedling EC",
            "Pest Check", "Pest Notes", "Leaf Cleaning", "Leaf Notes",
            "Notes", "Recorded By", "Created At"
        ]},
        row_marker=("Field Tasks", "Backend export test field"),
    )

    # Weekly
    test_export(
        token, "weekly", "/api/export/weekly",
        start=yesterday, end=tomorrow,
        expected_filename_prefix="Weekly",
        expected_sheets=["Weekly Checks"],
        expected_headers={"Weekly Checks": [
            "Date", "Meter Calibration", "Nutrition Qty OK", "Nutrition Notes",
            "Tank Filters Cleaned", "Notes", "Recorded By", "Created At"
        ]},
        row_marker=("Weekly Checks", "Backend export test weekly"),
    )

    # Monthly
    test_export(
        token, "monthly", "/api/export/monthly",
        start=yesterday, end=tomorrow,
        expected_filename_prefix="Monthly",
        expected_sheets=["Monthly Checks"],
        expected_headers={"Monthly Checks": [
            "Date", "Tanks Cleaned", "Salt Formation OK", "Salt Notes",
            "Solution A Qty", "Solution B Qty", "Solution C Qty", "Solutions Ordered",
            "Seeds Qty OK", "Seeds Ordered", "Notes", "Recorded By", "Created At"
        ]},
        row_marker=("Monthly Checks", "Backend export test monthly"),
    )

    print_summary()


def print_summary():
    print("\n===== SUMMARY =====")
    passed = sum(1 for _, p, _ in results if p)
    failed = [x for x in results if not x[1]]
    print(f"Total: {len(results)} | Passed: {passed} | Failed: {len(failed)}")
    if failed:
        print("\nFailed tests:")
        for name, _, msg in failed:
            print(f"  - {name} :: {msg}")


if __name__ == "__main__":
    main()
