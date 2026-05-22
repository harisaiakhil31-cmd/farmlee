"""Farmlee Manager Backend v2 - restructured per user requirements.

Sections:
- Tank readings (4 tanks, AM/Evening, pH/EC/Temp w/ target ranges)
- Environment readings (Morning/Afternoon/Evening temp+humidity, daily average)
- Field tasks (seedling watering, pest, leaf cleaning)
- Weekly checks (calibration, nutrition, filters)
- Monthly checks (tanks cleaning, salt, A/B/C, seeds)
- Custom reminders (push + email)
- Weekly report (aggregated data)
- Audit log (login history) - admin only, gated by extra OTP
- Crops library
"""
from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os, logging, secrets, bcrypt, jwt, uuid, smtplib, ssl
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from pathlib import Path
from pydantic import BaseModel, EmailStr
from typing import List, Optional
from datetime import datetime, timezone, timedelta
from sendgrid import SendGridAPIClient
from sendgrid.helpers.mail import Mail
import resend
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from openpyxl import Workbook
from io import BytesIO
from fastapi.responses import StreamingResponse

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

MONGO_URL = os.environ['MONGO_URL']
DB_NAME = os.environ['DB_NAME']
JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALG = os.environ.get('JWT_ALGORITHM', 'HS256')
SENDGRID_KEY = os.environ.get('SENDGRID_API_KEY', '')
RESEND_KEY = os.environ.get('RESEND_API_KEY', '')
GMAIL_USER = os.environ.get('GMAIL_USER', '')
GMAIL_APP_PASSWORD = os.environ.get('GMAIL_APP_PASSWORD', '').replace(' ', '')
SENDER_EMAIL = os.environ.get('SENDER_EMAIL', 'Farmlee Manager <noreply@farmlee.app>')
SENDER_NAME = os.environ.get('SENDER_NAME', 'Farmlee Manager')
ADMIN_EMAIL = os.environ['ADMIN_EMAIL']
ADMIN_PASSWORD = os.environ['ADMIN_PASSWORD']
MAX_USERS = int(os.environ.get('MAX_USERS', 3))

if RESEND_KEY:
    resend.api_key = RESEND_KEY

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]
app = FastAPI(title="Farmlee Manager API")
api = APIRouter(prefix="/api")
bearer = HTTPBearer(auto_error=False)
scheduler = AsyncIOScheduler()


def now_utc(): return datetime.now(timezone.utc)
def hash_pwd(p): return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()
def verify_pwd(p, h):
    try: return bcrypt.checkpw(p.encode(), h.encode())
    except: return False
def make_token(uid, email):
    return jwt.encode({"sub": uid, "email": email, "exp": now_utc() + timedelta(days=30), "iat": now_utc()}, JWT_SECRET, algorithm=JWT_ALG)

def _send_via_gmail(to, subject, html):
    """Send email via Gmail SMTP (sync)."""
    msg = MIMEMultipart("alternative")
    msg["Subject"] = subject
    msg["From"] = SENDER_EMAIL or GMAIL_USER
    msg["To"] = to
    msg.attach(MIMEText(html, "html"))
    ctx = ssl.create_default_context()
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ctx, timeout=15) as s:
        s.login(GMAIL_USER, GMAIL_APP_PASSWORD)
        s.sendmail(GMAIL_USER, [to], msg.as_string())

def send_email(to, subject, html):
    """Try Gmail SMTP -> Resend -> SendGrid -> dev log."""
    if GMAIL_USER and GMAIL_APP_PASSWORD:
        try:
            _send_via_gmail(to, subject, html)
            logger.info(f"Gmail OK -> {to} | {subject}")
            return True
        except Exception as e:
            logger.error(f"Gmail SMTP: {e}")
            # fall through
    if RESEND_KEY:
        try:
            resend.Emails.send({"from": SENDER_EMAIL, "to": [to], "subject": subject, "html": html})
            return True
        except Exception as e:
            logger.error(f"Resend: {e}")
    if SENDGRID_KEY:
        try:
            SendGridAPIClient(SENDGRID_KEY).send(Mail(from_email=SENDER_EMAIL, to_emails=to, subject=subject, html_content=html))
            return True
        except Exception as e:
            logger.error(f"SendGrid: {e}")
            return False
    logger.warning(f"[DEV-EMAIL] To={to} | {subject}\n{html}")
    return True

async def current_user(creds: HTTPAuthorizationCredentials = Depends(bearer)):
    if not creds: raise HTTPException(401, "Missing token")
    try: payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.PyJWTError: raise HTTPException(401, "Invalid token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user: raise HTTPException(401, "User not found")
    return user

async def admin_required(user=Depends(current_user)):
    if user.get("role") != "admin": raise HTTPException(403, "Admin only")
    return user


# ---------- MODELS ----------
class LoginIn(BaseModel):
    email: EmailStr; password: str; device: Optional[str] = "unknown"

class VerifyOTPIn(BaseModel):
    email: EmailStr; otp: str

class InviteUserIn(BaseModel):
    email: EmailStr; name: str; password: str

class TankReadingIn(BaseModel):
    check_date: str         # YYYY-MM-DD
    check_time: str         # HH:MM
    session: str            # "morning" | "evening"
    tank_id: int            # 1..4
    ph_actual: float
    ph_target_min: float
    ph_target_max: float
    ec_actual: float
    ec_target_min: float
    ec_target_max: float
    temp_actual: float
    temp_target_min: float
    temp_target_max: float
    notes: Optional[str] = ""

class EnvironmentReadingIn(BaseModel):
    check_date: str
    check_time: str
    session: str            # "morning" | "afternoon" | "evening"
    temperature: float
    humidity: float
    notes: Optional[str] = ""

class FieldTaskIn(BaseModel):
    check_date: str
    seedling_watered: bool = False
    seedling_ph: Optional[float] = None
    seedling_ec: Optional[float] = None
    pest_check_done: bool = False
    pest_notes: Optional[str] = ""
    leaf_cleaning_done: bool = False
    leaf_notes: Optional[str] = ""
    notes: Optional[str] = ""

class WeeklyCheckIn(BaseModel):
    check_date: str
    meter_calibration_done: bool = False
    nutrition_quantity_ok: bool = False
    nutrition_notes: str = ""
    tank_filters_cleaned: bool = False
    notes: str = ""

class MonthlyCheckIn(BaseModel):
    check_date: str
    tanks_cleaned: bool = False
    salt_formation_ok: bool = False
    salt_notes: str = ""
    solution_a_qty: Optional[float] = None
    solution_b_qty: Optional[float] = None
    solution_c_qty: Optional[float] = None
    solutions_ordered: bool = False
    seeds_qty_ok: bool = False
    seeds_ordered: bool = False
    notes: str = ""

class CropStage(BaseModel):
    name: str; ph_min: float; ph_max: float; ec_min: float; ec_max: float
    temp_min: float; temp_max: float; humidity_min: float; humidity_max: float
    duration_days: int; notes: str = ""

class CropIn(BaseModel):
    name: str; description: str = ""; stages: List[CropStage]

class ReminderIn(BaseModel):
    title: str; description: str = ""; remind_at: datetime
    notify_email: bool = True; notify_push: bool = True

class AdminOtpRequestIn(BaseModel):
    pass

# ---------- SEEDLING MANAGEMENT MODELS ----------
class SeedlingStageIn(BaseModel):
    stage_name: str
    order_index: int = 0
    duration_days: int
    ec_min: float
    ec_max: float
    water_temp_min: Optional[float] = None
    water_temp_max: Optional[float] = None
    notes: str = ""

class SeedlingTypeIn(BaseModel):
    name: str
    description: str = ""
    total_days_to_tower: int
    stages: List[SeedlingStageIn]
    notes: str = ""

class SeedlingBatchIn(BaseModel):
    seedling_type_id: str
    batch_name: str
    quantity: int = 0
    sown_date: str            # YYYY-MM-DD
    expected_transplant_date: Optional[str] = None
    tray_location: str = ""
    notes: str = ""

class SeedlingBatchUpdate(BaseModel):
    batch_name: Optional[str] = None
    quantity: Optional[int] = None
    expected_transplant_date: Optional[str] = None
    tray_location: Optional[str] = None
    notes: Optional[str] = None
    current_stage_index: Optional[int] = None

class SeedlingTransplantIn(BaseModel):
    actual_transplant_date: str
    tower_destination: str = ""
    notes: str = ""

class SeedlingWateringIn(BaseModel):
    batch_id: str
    log_date: str             # YYYY-MM-DD
    session: str              # "morning" | "evening"
    watered: bool = True
    ec_before: Optional[float] = None
    water_temp_before: Optional[float] = None
    ec_target_min: Optional[float] = None
    ec_target_max: Optional[float] = None
    ec_actual: Optional[float] = None
    notes: str = ""

class AdminOtpVerifyIn(BaseModel):
    otp: str

class ChangePasswordIn(BaseModel):
    current_password: str
    new_password: str

class ForgotPasswordIn(BaseModel):
    email: EmailStr

class ResetPasswordIn(BaseModel):
    email: EmailStr
    otp: str
    new_password: str


import re
def validate_password_strength(pwd: str):
    """Raise HTTPException if password doesn't meet policy."""
    if len(pwd) < 8:
        raise HTTPException(400, "Password must be at least 8 characters")
    if not re.search(r"[A-Z]", pwd):
        raise HTTPException(400, "Password must include an uppercase letter")
    if not re.search(r"[a-z]", pwd):
        raise HTTPException(400, "Password must include a lowercase letter")
    if not re.search(r"\d", pwd):
        raise HTTPException(400, "Password must include a number")
    if not re.search(r"[^A-Za-z0-9]", pwd):
        raise HTTPException(400, "Password must include a special character")


# ---------- AUTH + AUDIT ----------
@api.post("/auth/login")
async def login(body: LoginIn, request: Request):
    user = await db.users.find_one({"email": body.email.lower()})
    if not user or not verify_pwd(body.password, user["password_hash"]):
        raise HTTPException(401, "Invalid email or password")
    otp = f"{secrets.randbelow(1000000):06d}"
    await db.otp_codes.delete_many({"email": body.email.lower(), "purpose": "login"})
    await db.otp_codes.insert_one({"email": body.email.lower(), "otp": otp, "purpose": "login",
        "expires_at": now_utc() + timedelta(minutes=10), "created_at": now_utc()})
    html = f"<div style='font-family:Arial;padding:24px;background:#F9F8F6'><h2 style='color:#1B2E1C'>Farmlee Manager Login Code</h2><p>Hi {user.get('name','')}, your sign-in code:</p><div style='font-size:36px;letter-spacing:8px;font-weight:700;color:#4A5D23;background:#fff;padding:20px;text-align:center;border-radius:12px'>{otp}</div><p style='color:#888;font-size:12px'>Expires in 10 minutes.</p></div>"
    send_email(body.email, "Farmlee Manager verification code", html)
    # Store pending login (for audit on success)
    await db.pending_logins.delete_many({"email": body.email.lower()})
    await db.pending_logins.insert_one({
        "email": body.email.lower(),
        "device": body.device or "unknown",
        "ip": request.client.host if request.client else "unknown",
        "user_agent": request.headers.get("user-agent", "unknown")[:300],
        "created_at": now_utc(),
    })
    resp = {"message": "OTP sent to your email", "email": body.email.lower()}
    if not SENDGRID_KEY and not RESEND_KEY and not GMAIL_USER: resp["dev_otp"] = otp
    return resp

@api.post("/auth/verify-otp")
async def verify_otp(body: VerifyOTPIn, request: Request):
    rec = await db.otp_codes.find_one({"email": body.email.lower(), "otp": body.otp, "purpose": "login"})
    if not rec: raise HTTPException(401, "Invalid OTP")
    exp = rec["expires_at"]
    if exp.tzinfo is None: exp = exp.replace(tzinfo=timezone.utc)
    if exp < now_utc(): raise HTTPException(401, "OTP expired")
    user = await db.users.find_one({"email": body.email.lower()}, {"_id": 0, "password_hash": 0})
    if not user: raise HTTPException(404, "User not found")
    await db.otp_codes.delete_many({"email": body.email.lower(), "purpose": "login"})
    token = make_token(user["id"], user["email"])
    # write audit log
    pending = await db.pending_logins.find_one({"email": body.email.lower()})
    await db.audit_logs.insert_one({
        "id": str(uuid.uuid4()),
        "user_id": user["id"], "user_name": user["name"], "email": user["email"],
        "device": (pending or {}).get("device", "unknown"),
        "ip": (pending or {}).get("ip", request.client.host if request.client else "unknown"),
        "user_agent": (pending or {}).get("user_agent", request.headers.get("user-agent", "unknown")[:300]),
        "logged_in_at": now_utc(), "event": "login_success",
    })
    await db.pending_logins.delete_many({"email": body.email.lower()})
    return {"token": token, "user": user}

@api.get("/auth/me")
async def me(user=Depends(current_user)): return user


# ---------- USERS ----------
@api.get("/users")
async def list_users(admin=Depends(admin_required)):
    return await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(100)

@api.post("/users/invite")
async def invite_user(body: InviteUserIn, admin=Depends(admin_required)):
    if await db.users.count_documents({}) >= MAX_USERS: raise HTTPException(400, f"Max {MAX_USERS} users reached")
    if await db.users.find_one({"email": body.email.lower()}): raise HTTPException(400, "Email already exists")
    u = {"id": str(uuid.uuid4()), "email": body.email.lower(), "name": body.name,
         "password_hash": hash_pwd(body.password), "role": "member", "created_at": now_utc()}
    await db.users.insert_one(u)
    send_email(body.email, "Welcome to Farmlee Manager",
        f"<div style='font-family:Arial;padding:24px'><h2>Welcome to Farmlee Manager</h2><p>Email: {body.email}<br/>Temporary password: {body.password}</p></div>")
    return {"id": u["id"], "email": u["email"], "name": u["name"]}

@api.delete("/users/{uid}")
async def delete_user(uid: str, admin=Depends(admin_required)):
    if uid == admin["id"]: raise HTTPException(400, "Cannot delete yourself")
    r = await db.users.delete_one({"id": uid, "role": {"$ne": "admin"}})
    if r.deleted_count == 0: raise HTTPException(404, "User not found or is admin")
    return {"ok": True}


# ---------- TANK READINGS ----------
@api.post("/tanks/reading")
async def create_tank(body: TankReadingIn, user=Depends(current_user)):
    if body.tank_id < 1 or body.tank_id > 4: raise HTTPException(400, "tank_id must be 1..4")
    if body.session not in ("morning", "evening"): raise HTTPException(400, "session must be morning|evening")
    d = body.model_dump()
    d.update({"id": str(uuid.uuid4()), "user_id": user["id"], "user_name": user["name"], "created_at": now_utc()})
    await db.tank_readings.insert_one(d.copy()); d.pop("_id", None); return d

@api.get("/tanks/reading")
async def list_tanks(date: Optional[str] = None, user=Depends(current_user)):
    q = {}
    if date: q["check_date"] = date
    return await db.tank_readings.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


# ---------- ENVIRONMENT ----------
@api.post("/environment/reading")
async def create_env(body: EnvironmentReadingIn, user=Depends(current_user)):
    if body.session not in ("morning", "afternoon", "evening"): raise HTTPException(400, "Invalid session")
    d = body.model_dump()
    d.update({"id": str(uuid.uuid4()), "user_id": user["id"], "user_name": user["name"], "created_at": now_utc()})
    await db.environment_readings.insert_one(d.copy()); d.pop("_id", None); return d

@api.get("/environment/reading")
async def list_env(date: Optional[str] = None, user=Depends(current_user)):
    q = {}
    if date: q["check_date"] = date
    items = await db.environment_readings.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    avg = None
    if items:
        temps = [i["temperature"] for i in items]; hums = [i["humidity"] for i in items]
        avg = {"temperature": round(sum(temps)/len(temps), 2), "humidity": round(sum(hums)/len(hums), 2), "count": len(items)}
    return {"items": items, "average": avg}


# ---------- FIELD TASKS ----------
@api.post("/field/tasks")
async def create_field(body: FieldTaskIn, user=Depends(current_user)):
    d = body.model_dump()
    d.update({"id": str(uuid.uuid4()), "user_id": user["id"], "user_name": user["name"], "created_at": now_utc()})
    await db.field_tasks.insert_one(d.copy()); d.pop("_id", None); return d

@api.get("/field/tasks")
async def list_field(date: Optional[str] = None, user=Depends(current_user)):
    q = {}
    if date: q["check_date"] = date
    return await db.field_tasks.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


# ---------- WEEKLY / MONTHLY CHECKS ----------
@api.post("/checks/weekly")
async def create_weekly(body: WeeklyCheckIn, user=Depends(current_user)):
    d = body.model_dump()
    d.update({"id": str(uuid.uuid4()), "user_id": user["id"], "user_name": user["name"], "created_at": now_utc()})
    await db.weekly_checks.insert_one(d.copy()); d.pop("_id", None); return d

@api.get("/checks/weekly")
async def list_weekly(user=Depends(current_user)):
    return await db.weekly_checks.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)

@api.post("/checks/monthly")
async def create_monthly(body: MonthlyCheckIn, user=Depends(current_user)):
    d = body.model_dump()
    d.update({"id": str(uuid.uuid4()), "user_id": user["id"], "user_name": user["name"], "created_at": now_utc()})
    await db.monthly_checks.insert_one(d.copy()); d.pop("_id", None); return d

@api.get("/checks/monthly")
async def list_monthly(user=Depends(current_user)):
    return await db.monthly_checks.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)


# ---------- WEEKLY REPORT ----------
@api.get("/report/weekly")
async def weekly_report(start: str, user=Depends(current_user)):
    """Return aggregated data for 7 days starting `start` (YYYY-MM-DD)."""
    start_d = datetime.fromisoformat(start).date()
    end_d = start_d + timedelta(days=6)
    s, e = start_d.isoformat(), end_d.isoformat()
    tanks = await db.tank_readings.find({"check_date": {"$gte": s, "$lte": e}}, {"_id": 0}).to_list(2000)
    env = await db.environment_readings.find({"check_date": {"$gte": s, "$lte": e}}, {"_id": 0}).to_list(2000)
    field = await db.field_tasks.find({"check_date": {"$gte": s, "$lte": e}}, {"_id": 0}).to_list(2000)
    weekly = await db.weekly_checks.find({"check_date": {"$gte": s, "$lte": e}}, {"_id": 0}).to_list(2000)
    # by day
    days = []
    for i in range(7):
        d = (start_d + timedelta(days=i)).isoformat()
        d_tanks = [t for t in tanks if t["check_date"] == d]
        d_env = [e for e in env if e["check_date"] == d]
        avg = None
        if d_env:
            avg = {"temperature": round(sum(x["temperature"] for x in d_env)/len(d_env), 1),
                   "humidity": round(sum(x["humidity"] for x in d_env)/len(d_env), 1)}
        tank_avg = None
        if d_tanks:
            tank_avg = {
                "ph": round(sum(t["ph_actual"] for t in d_tanks)/len(d_tanks), 2),
                "ec": round(sum(t["ec_actual"] for t in d_tanks)/len(d_tanks), 2),
                "temp": round(sum(t["temp_actual"] for t in d_tanks)/len(d_tanks), 1),
            }
        days.append({"date": d, "tank_readings_count": len(d_tanks), "tank_avg": tank_avg,
                     "env_readings_count": len(d_env), "env_avg": avg,
                     "field_tasks_count": len([f for f in field if f["check_date"] == d])})
    return {"start": s, "end": e, "days": days, "totals": {
        "tank_readings": len(tanks), "environment_readings": len(env),
        "field_tasks": len(field), "weekly_checks": len(weekly),
    }}


# ---------- CALENDAR ----------
@api.get("/calendar")
async def calendar(start: str, end: str, user=Depends(current_user)):
    tanks = await db.tank_readings.find({"check_date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(2000)
    env = await db.environment_readings.find({"check_date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(2000)
    field = await db.field_tasks.find({"check_date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(2000)
    weekly = await db.weekly_checks.find({"check_date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(500)
    monthly = await db.monthly_checks.find({"check_date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(500)
    reminders = await db.reminders.find({"remind_at": {
        "$gte": datetime.fromisoformat(start).replace(tzinfo=timezone.utc),
        "$lte": datetime.fromisoformat(end).replace(tzinfo=timezone.utc) + timedelta(days=1)}}, {"_id": 0}).to_list(500)
    return {"tanks": tanks, "environment": env, "field": field, "weekly": weekly, "monthly": monthly, "reminders": reminders}

@api.get("/calendar/day/{day}")
async def day_detail(day: str, user=Depends(current_user)):
    tanks = await db.tank_readings.find({"check_date": day}, {"_id": 0}).to_list(200)
    env = await db.environment_readings.find({"check_date": day}, {"_id": 0}).to_list(200)
    field = await db.field_tasks.find({"check_date": day}, {"_id": 0}).to_list(200)
    weekly = await db.weekly_checks.find({"check_date": day}, {"_id": 0}).to_list(50)
    monthly = await db.monthly_checks.find({"check_date": day}, {"_id": 0}).to_list(50)
    env_avg = None
    if env:
        env_avg = {"temperature": round(sum(x["temperature"] for x in env)/len(env), 1),
                   "humidity": round(sum(x["humidity"] for x in env)/len(env), 1)}
    today = now_utc().strftime("%Y-%m-%d")
    return {"date": day, "is_future": day > today, "is_today": day == today,
            "tanks": tanks, "environment": env, "env_avg": env_avg,
            "field": field, "weekly": weekly, "monthly": monthly}


# ---------- CROPS ----------
@api.get("/crops")
async def list_crops(user=Depends(current_user)):
    return await db.crops.find({}, {"_id": 0}).sort("name", 1).to_list(100)

@api.post("/crops")
async def create_crop(body: CropIn, user=Depends(current_user)):
    c = body.model_dump(); c.update({"id": str(uuid.uuid4()), "created_at": now_utc(), "created_by": user["id"]})
    await db.crops.insert_one(c.copy()); c.pop("_id", None); return c

@api.put("/crops/{cid}")
async def update_crop(cid: str, body: CropIn, user=Depends(current_user)):
    u = body.model_dump(); u["updated_at"] = now_utc()
    r = await db.crops.update_one({"id": cid}, {"$set": u})
    if r.matched_count == 0: raise HTTPException(404, "Not found")
    return await db.crops.find_one({"id": cid}, {"_id": 0})

@api.delete("/crops/{cid}")
async def delete_crop(cid: str, user=Depends(current_user)):
    await db.crops.delete_one({"id": cid}); return {"ok": True}


# ---------- REMINDERS ----------
@api.get("/reminders")
async def list_reminders(user=Depends(current_user)):
    return await db.reminders.find({"user_id": user["id"]}, {"_id": 0}).sort("remind_at", 1).to_list(500)

@api.post("/reminders")
async def create_reminder(body: ReminderIn, user=Depends(current_user)):
    r = body.model_dump()
    if r["remind_at"].tzinfo is None: r["remind_at"] = r["remind_at"].replace(tzinfo=timezone.utc)
    r.update({"id": str(uuid.uuid4()), "user_id": user["id"], "user_email": user["email"],
              "fired": False, "created_at": now_utc()})
    await db.reminders.insert_one(r.copy()); r.pop("_id", None); return r

@api.delete("/reminders/{rid}")
async def delete_reminder(rid: str, user=Depends(current_user)):
    await db.reminders.delete_one({"id": rid, "user_id": user["id"]}); return {"ok": True}


# ---------- ADMIN AUDIT LOG (gated by extra OTP) ----------
@api.post("/admin/audit/request-otp")
async def admin_request_otp(admin=Depends(admin_required)):
    otp = f"{secrets.randbelow(1000000):06d}"
    await db.otp_codes.delete_many({"email": admin["email"], "purpose": "audit"})
    await db.otp_codes.insert_one({"email": admin["email"], "otp": otp, "purpose": "audit",
        "expires_at": now_utc() + timedelta(minutes=10), "created_at": now_utc()})
    send_email(admin["email"], "Farmlee Manager audit access code",
        f"<div style='font-family:Arial;padding:24px;background:#F9F8F6'><h2 style='color:#1B2E1C'>Audit Log Access</h2><p>Use this code to view the audit log:</p><div style='font-size:36px;letter-spacing:8px;font-weight:700;color:#CC7753;background:#fff;padding:20px;text-align:center;border-radius:12px'>{otp}</div></div>")
    resp = {"message": "OTP sent"}
    if not SENDGRID_KEY and not RESEND_KEY and not GMAIL_USER: resp["dev_otp"] = otp
    return resp

@api.post("/admin/audit/verify")
async def admin_verify_audit(body: AdminOtpVerifyIn, admin=Depends(admin_required)):
    rec = await db.otp_codes.find_one({"email": admin["email"], "otp": body.otp, "purpose": "audit"})
    if not rec: raise HTTPException(401, "Invalid OTP")
    exp = rec["expires_at"]
    if exp.tzinfo is None: exp = exp.replace(tzinfo=timezone.utc)
    if exp < now_utc(): raise HTTPException(401, "OTP expired")
    await db.otp_codes.delete_many({"email": admin["email"], "purpose": "audit"})
    logs = await db.audit_logs.find({}, {"_id": 0}).sort("logged_in_at", -1).limit(200).to_list(200)
    return {"logs": logs}


# ---------- PASSWORD CHANGE / RESET ----------
@api.post("/auth/change-password")
async def change_password(body: ChangePasswordIn, user=Depends(current_user)):
    full = await db.users.find_one({"id": user["id"]})
    if not full or not verify_pwd(body.current_password, full["password_hash"]):
        raise HTTPException(401, "Current password is incorrect")
    if body.current_password == body.new_password:
        raise HTTPException(400, "New password must be different from the current password")
    validate_password_strength(body.new_password)
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "password_hash": hash_pwd(body.new_password),
        "password_changed_at": now_utc(),
    }})
    return {"ok": True, "message": "Password updated"}

@api.post("/auth/forgot-password")
async def forgot_password(body: ForgotPasswordIn):
    """Always returns success to avoid leaking which emails exist."""
    user = await db.users.find_one({"email": body.email.lower()})
    resp = {"message": "If this email is registered, a reset code has been sent."}
    if not user:
        return resp
    otp = f"{secrets.randbelow(1000000):06d}"
    await db.otp_codes.delete_many({"email": body.email.lower(), "purpose": "reset"})
    await db.otp_codes.insert_one({"email": body.email.lower(), "otp": otp, "purpose": "reset",
        "expires_at": now_utc() + timedelta(minutes=15), "created_at": now_utc()})
    html = f"<div style='font-family:Arial;padding:24px;background:#F9F8F6'><h2 style='color:#1B2E1C'>Password reset code</h2><p>Hi {user.get('name','')}, use this code to reset your Farmlee Manager password:</p><div style='font-size:36px;letter-spacing:8px;font-weight:700;color:#CC7753;background:#fff;padding:20px;text-align:center;border-radius:12px'>{otp}</div><p style='color:#888;font-size:12px'>Expires in 15 minutes. If you didn't request this, you can ignore this email.</p></div>"
    send_email(body.email, "Farmlee Manager password reset code", html)
    if not SENDGRID_KEY and not RESEND_KEY and not GMAIL_USER: resp["dev_otp"] = otp
    return resp

@api.post("/auth/reset-password")
async def reset_password(body: ResetPasswordIn):
    rec = await db.otp_codes.find_one({"email": body.email.lower(), "otp": body.otp, "purpose": "reset"})
    if not rec: raise HTTPException(401, "Invalid or expired code")
    exp = rec["expires_at"]
    if exp.tzinfo is None: exp = exp.replace(tzinfo=timezone.utc)
    if exp < now_utc(): raise HTTPException(401, "Reset code expired")
    user = await db.users.find_one({"email": body.email.lower()})
    if not user: raise HTTPException(404, "User not found")
    validate_password_strength(body.new_password)
    await db.otp_codes.delete_many({"email": body.email.lower(), "purpose": "reset"})
    await db.users.update_one({"id": user["id"]}, {"$set": {
        "password_hash": hash_pwd(body.new_password),
        "password_changed_at": now_utc(),
    }})
    return {"ok": True, "message": "Password has been reset. Please sign in."}


# ---------- SEEDLING MANAGEMENT ----------
@api.get("/seedlings/types")
async def list_seedling_types(user=Depends(current_user)):
    return await db.seedling_types.find({}, {"_id": 0}).sort("name", 1).to_list(200)

@api.post("/seedlings/types")
async def create_seedling_type(body: SeedlingTypeIn, user=Depends(current_user)):
    t = body.model_dump()
    # ensure stages ordered
    t["stages"] = sorted(t["stages"], key=lambda s: s.get("order_index", 0))
    t.update({"id": str(uuid.uuid4()), "created_at": now_utc(), "created_by": user["id"]})
    await db.seedling_types.insert_one(t.copy()); t.pop("_id", None); return t

@api.put("/seedlings/types/{tid}")
async def update_seedling_type(tid: str, body: SeedlingTypeIn, user=Depends(current_user)):
    u = body.model_dump()
    u["stages"] = sorted(u["stages"], key=lambda s: s.get("order_index", 0))
    u["updated_at"] = now_utc()
    r = await db.seedling_types.update_one({"id": tid}, {"$set": u})
    if r.matched_count == 0: raise HTTPException(404, "Seedling type not found")
    return await db.seedling_types.find_one({"id": tid}, {"_id": 0})

@api.delete("/seedlings/types/{tid}")
async def delete_seedling_type(tid: str, user=Depends(current_user)):
    # block delete if active batches exist
    active = await db.seedling_batches.count_documents({"seedling_type_id": tid, "status": "active"})
    if active: raise HTTPException(400, f"Cannot delete: {active} active batch(es) of this type")
    await db.seedling_types.delete_one({"id": tid}); return {"ok": True}

# ----- BATCHES -----
def _compute_stage_progress(batch, sed_type):
    """Given a batch + its seedling type, compute current stage index, day-in-stage, days-since-sown."""
    try:
        sown = datetime.strptime(batch["sown_date"], "%Y-%m-%d").date()
    except Exception:
        return {"current_stage_index": 0, "day_in_stage": 0, "days_since_sown": 0, "current_stage": None}
    today_d = datetime.now(timezone.utc).date()
    days_since = (today_d - sown).days
    if days_since < 0: days_since = 0
    stages = sed_type.get("stages", []) if sed_type else []
    elapsed = 0
    idx = 0
    for i, s in enumerate(stages):
        d = max(1, int(s.get("duration_days", 0) or 0))
        if days_since < elapsed + d:
            idx = i
            return {"current_stage_index": i, "day_in_stage": days_since - elapsed + 1,
                    "days_since_sown": days_since, "current_stage": s}
        elapsed += d
        idx = i
    # past last stage -> ready to transplant
    last = stages[-1] if stages else None
    return {"current_stage_index": idx, "day_in_stage": (days_since - (elapsed - (last.get("duration_days",0) if last else 0))) if last else 0,
            "days_since_sown": days_since, "current_stage": last, "ready_to_transplant": True}

@api.get("/seedlings/batches")
async def list_seedling_batches(status: Optional[str] = None, user=Depends(current_user)):
    q = {}
    if status: q["status"] = status
    batches = await db.seedling_batches.find(q, {"_id": 0}).sort("sown_date", -1).to_list(500)
    types = {t["id"]: t for t in await db.seedling_types.find({}, {"_id": 0}).to_list(500)}
    for b in batches:
        t = types.get(b.get("seedling_type_id"))
        b["seedling_type_name"] = t["name"] if t else "Unknown"
        b["progress"] = _compute_stage_progress(b, t) if t else {}
    return batches

@api.post("/seedlings/batches")
async def create_seedling_batch(body: SeedlingBatchIn, user=Depends(current_user)):
    t = await db.seedling_types.find_one({"id": body.seedling_type_id}, {"_id": 0})
    if not t: raise HTTPException(404, "Seedling type not found")
    b = body.model_dump()
    if not b.get("expected_transplant_date"):
        try:
            sown = datetime.strptime(b["sown_date"], "%Y-%m-%d").date()
            b["expected_transplant_date"] = (sown + timedelta(days=t.get("total_days_to_tower", 21))).isoformat()
        except Exception:
            pass
    b.update({"id": str(uuid.uuid4()), "status": "active",
              "current_stage_index": 0, "actual_transplant_date": None, "tower_destination": None,
              "created_at": now_utc(), "created_by": user["id"], "user_name": user["name"]})
    await db.seedling_batches.insert_one(b.copy()); b.pop("_id", None); return b

@api.get("/seedlings/batches/{bid}")
async def get_seedling_batch(bid: str, user=Depends(current_user)):
    b = await db.seedling_batches.find_one({"id": bid}, {"_id": 0})
    if not b: raise HTTPException(404, "Batch not found")
    t = await db.seedling_types.find_one({"id": b["seedling_type_id"]}, {"_id": 0})
    b["seedling_type"] = t
    b["seedling_type_name"] = t["name"] if t else "Unknown"
    b["progress"] = _compute_stage_progress(b, t) if t else {}
    b["watering_logs"] = await db.seedling_watering.find({"batch_id": bid}, {"_id": 0}).sort("log_date", -1).to_list(500)
    return b

@api.put("/seedlings/batches/{bid}")
async def update_seedling_batch(bid: str, body: SeedlingBatchUpdate, user=Depends(current_user)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    upd["updated_at"] = now_utc()
    r = await db.seedling_batches.update_one({"id": bid}, {"$set": upd})
    if r.matched_count == 0: raise HTTPException(404, "Batch not found")
    return await db.seedling_batches.find_one({"id": bid}, {"_id": 0})

@api.post("/seedlings/batches/{bid}/transplant")
async def transplant_batch(bid: str, body: SeedlingTransplantIn, user=Depends(current_user)):
    r = await db.seedling_batches.update_one({"id": bid}, {"$set": {
        "status": "transplanted",
        "actual_transplant_date": body.actual_transplant_date,
        "tower_destination": body.tower_destination,
        "transplant_notes": body.notes,
        "transplanted_at": now_utc(),
        "transplanted_by": user["id"],
    }})
    if r.matched_count == 0: raise HTTPException(404, "Batch not found")
    return {"ok": True}

@api.delete("/seedlings/batches/{bid}")
async def delete_seedling_batch(bid: str, user=Depends(current_user)):
    await db.seedling_watering.delete_many({"batch_id": bid})
    await db.seedling_batches.delete_one({"id": bid})
    return {"ok": True}

# ----- WATERING LOGS -----
@api.post("/seedlings/watering")
async def create_watering_log(body: SeedlingWateringIn, user=Depends(current_user)):
    if body.session not in ("morning", "evening"):
        raise HTTPException(400, "session must be morning|evening")
    if not await db.seedling_batches.find_one({"id": body.batch_id}):
        raise HTTPException(404, "Batch not found")
    d = body.model_dump()
    d.update({"id": str(uuid.uuid4()), "recorded_by": user["id"], "user_name": user["name"], "created_at": now_utc()})
    await db.seedling_watering.insert_one(d.copy()); d.pop("_id", None); return d

@api.get("/seedlings/watering")
async def list_watering(batch_id: Optional[str] = None, date: Optional[str] = None, user=Depends(current_user)):
    q = {}
    if batch_id: q["batch_id"] = batch_id
    if date: q["log_date"] = date
    return await db.seedling_watering.find(q, {"_id": 0}).sort("log_date", -1).to_list(1000)

@api.delete("/seedlings/watering/{wid}")
async def delete_watering_log(wid: str, user=Depends(current_user)):
    await db.seedling_watering.delete_one({"id": wid}); return {"ok": True}


# ---------- DASHBOARD ----------
@api.get("/dashboard")
async def dashboard(user=Depends(current_user)):
    today = now_utc().strftime("%Y-%m-%d")
    tanks = await db.tank_readings.count_documents({"check_date": today})
    env = await db.environment_readings.find({"check_date": today}, {"_id": 0}).to_list(20)
    field = await db.field_tasks.count_documents({"check_date": today})
    env_avg = None
    if env:
        env_avg = {"temperature": round(sum(x["temperature"] for x in env)/len(env), 1),
                   "humidity": round(sum(x["humidity"] for x in env)/len(env), 1),
                   "sessions": [e["session"] for e in env]}
    upcoming = await db.reminders.find({"user_id": user["id"], "remind_at": {"$gte": now_utc()}, "fired": False},
        {"_id": 0}).sort("remind_at", 1).limit(5).to_list(5)
    total_users = await db.users.count_documents({})
    return {"today": today, "tank_readings_today": tanks, "env_today_count": len(env),
            "env_today_avg": env_avg, "field_tasks_today": field,
            "upcoming_reminders": upcoming, "total_users": total_users, "max_users": MAX_USERS}


# ---------- SCHEDULER ----------
async def reminder_worker():
    try:
        async for r in db.reminders.find({"fired": False, "remind_at": {"$lte": now_utc()}}):
            try:
                if r.get("notify_email", True):
                    send_email(r["user_email"], f"Reminder: {r.get('title','')}",
                        f"<div style='font-family:Arial;padding:24px;background:#F9F8F6'><h2 style='color:#1B2E1C'>🌱 Farmlee Manager Reminder</h2><h3 style='color:#4A5D23'>{r.get('title','')}</h3><p>{r.get('description','')}</p></div>")
                await db.reminders.update_one({"id": r["id"]}, {"$set": {"fired": True, "fired_at": now_utc()}})
            except Exception as e:
                logger.error(f"Fire: {e}")
    except Exception as e:
        logger.error(f"Worker: {e}")


# ---------- STARTUP ----------
@app.on_event("startup")
async def startup():
    if not await db.users.find_one({"email": ADMIN_EMAIL.lower()}):
        await db.users.insert_one({"id": str(uuid.uuid4()), "email": ADMIN_EMAIL.lower(), "name": "Admin",
            "password_hash": hash_pwd(ADMIN_PASSWORD), "role": "admin", "created_at": now_utc()})
        logger.info(f"Seeded admin {ADMIN_EMAIL}")
    if await db.crops.count_documents({}) == 0:
        examples = [
            {"id": str(uuid.uuid4()), "name": "Lettuce", "description": "Leafy green, fast-growing.",
             "stages": [{"name": "seedling", "ph_min": 5.8, "ph_max": 6.2, "ec_min": 0.4, "ec_max": 0.8, "temp_min": 18, "temp_max": 22, "humidity_min": 60, "humidity_max": 75, "duration_days": 14, "notes": ""},
                        {"name": "vegetative", "ph_min": 5.5, "ph_max": 6.5, "ec_min": 0.8, "ec_max": 1.2, "temp_min": 18, "temp_max": 24, "humidity_min": 50, "humidity_max": 70, "duration_days": 21, "notes": ""},
                        {"name": "harvest", "ph_min": 5.5, "ph_max": 6.5, "ec_min": 1.2, "ec_max": 1.8, "temp_min": 16, "temp_max": 22, "humidity_min": 50, "humidity_max": 65, "duration_days": 7, "notes": ""}],
             "created_at": now_utc()},
            {"id": str(uuid.uuid4()), "name": "Tomato", "description": "Fruiting crop.",
             "stages": [{"name": "seedling", "ph_min": 5.8, "ph_max": 6.2, "ec_min": 0.8, "ec_max": 1.2, "temp_min": 20, "temp_max": 25, "humidity_min": 65, "humidity_max": 75, "duration_days": 21, "notes": ""},
                        {"name": "flowering", "ph_min": 5.8, "ph_max": 6.5, "ec_min": 2.0, "ec_max": 3.0, "temp_min": 21, "temp_max": 26, "humidity_min": 55, "humidity_max": 65, "duration_days": 21, "notes": ""},
                        {"name": "fruiting", "ph_min": 6.0, "ph_max": 6.5, "ec_min": 2.5, "ec_max": 3.5, "temp_min": 22, "temp_max": 27, "humidity_min": 55, "humidity_max": 65, "duration_days": 60, "notes": ""}],
             "created_at": now_utc()},
            {"id": str(uuid.uuid4()), "name": "Basil", "description": "Aromatic herb.",
             "stages": [{"name": "seedling", "ph_min": 5.5, "ph_max": 6.5, "ec_min": 0.6, "ec_max": 1.0, "temp_min": 20, "temp_max": 25, "humidity_min": 60, "humidity_max": 75, "duration_days": 14, "notes": ""},
                        {"name": "vegetative", "ph_min": 5.5, "ph_max": 6.5, "ec_min": 1.0, "ec_max": 1.6, "temp_min": 20, "temp_max": 26, "humidity_min": 50, "humidity_max": 65, "duration_days": 35, "notes": ""}],
             "created_at": now_utc()},
        ]
        await db.crops.insert_many(examples); logger.info("Seeded crops")
    if await db.seedling_types.count_documents({}) == 0:
        seedling_examples = [
            {"id": str(uuid.uuid4()), "name": "Lettuce", "description": "Leafy green seedling.",
             "total_days_to_tower": 21,
             "stages": [
                 {"stage_name": "Germination", "order_index": 0, "duration_days": 4, "ec_min": 0.2, "ec_max": 0.4, "water_temp_min": 20, "water_temp_max": 24, "notes": "Keep moist"},
                 {"stage_name": "Cotyledon",   "order_index": 1, "duration_days": 5, "ec_min": 0.4, "ec_max": 0.8, "water_temp_min": 19, "water_temp_max": 23, "notes": "First leaves emerge"},
                 {"stage_name": "True Leaf",   "order_index": 2, "duration_days": 7, "ec_min": 0.8, "ec_max": 1.2, "water_temp_min": 18, "water_temp_max": 22, "notes": "Ready to harden off"},
                 {"stage_name": "Pre-Transplant", "order_index": 3, "duration_days": 5, "ec_min": 1.0, "ec_max": 1.4, "water_temp_min": 18, "water_temp_max": 22, "notes": "Harden before tower"},
             ], "notes": "", "created_at": now_utc()},
            {"id": str(uuid.uuid4()), "name": "Basil", "description": "Aromatic herb seedling.",
             "total_days_to_tower": 25,
             "stages": [
                 {"stage_name": "Germination",  "order_index": 0, "duration_days": 5, "ec_min": 0.2, "ec_max": 0.4, "water_temp_min": 22, "water_temp_max": 26, "notes": ""},
                 {"stage_name": "Cotyledon",    "order_index": 1, "duration_days": 6, "ec_min": 0.5, "ec_max": 0.9, "water_temp_min": 21, "water_temp_max": 25, "notes": ""},
                 {"stage_name": "True Leaf",    "order_index": 2, "duration_days": 8, "ec_min": 0.9, "ec_max": 1.4, "water_temp_min": 20, "water_temp_max": 24, "notes": ""},
                 {"stage_name": "Pre-Transplant","order_index": 3, "duration_days": 6, "ec_min": 1.2, "ec_max": 1.6, "water_temp_min": 20, "water_temp_max": 24, "notes": ""},
             ], "notes": "", "created_at": now_utc()},
            {"id": str(uuid.uuid4()), "name": "Tomato", "description": "Fruiting crop seedling.",
             "total_days_to_tower": 28,
             "stages": [
                 {"stage_name": "Germination", "order_index": 0, "duration_days": 6, "ec_min": 0.3, "ec_max": 0.5, "water_temp_min": 22, "water_temp_max": 26, "notes": ""},
                 {"stage_name": "Cotyledon",   "order_index": 1, "duration_days": 7, "ec_min": 0.6, "ec_max": 1.0, "water_temp_min": 21, "water_temp_max": 25, "notes": ""},
                 {"stage_name": "True Leaf",   "order_index": 2, "duration_days": 9, "ec_min": 1.0, "ec_max": 1.6, "water_temp_min": 20, "water_temp_max": 24, "notes": ""},
                 {"stage_name": "Pre-Transplant","order_index": 3, "duration_days": 6, "ec_min": 1.4, "ec_max": 2.0, "water_temp_min": 20, "water_temp_max": 24, "notes": ""},
             ], "notes": "", "created_at": now_utc()},
        ]
        await db.seedling_types.insert_many(seedling_examples); logger.info("Seeded seedling types")
    scheduler.add_job(reminder_worker, "interval", minutes=1, id="reminder_worker", replace_existing=True)
    if not scheduler.running: scheduler.start()


@app.on_event("shutdown")
async def shutdown():
    try: scheduler.shutdown(wait=False)
    except: pass
    client.close()


# ---------- EXCEL EXPORT ----------
def _autosize(ws):
    for col in ws.columns:
        try:
            letter = col[0].column_letter
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 40)
        except Exception:
            pass

def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

def _date_filter(start: Optional[str], end: Optional[str]):
    q = {}
    if start and end: q["check_date"] = {"$gte": start, "$lte": end}
    elif start: q["check_date"] = {"$gte": start}
    elif end: q["check_date"] = {"$lte": end}
    return q

def _fmt_dt(v):
    if isinstance(v, datetime):
        try: return v.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        except: return v.isoformat()
    return v

@api.get("/export/tanks")
async def export_tanks(start: Optional[str] = None, end: Optional[str] = None, user=Depends(current_user)):
    rows = await db.tank_readings.find(_date_filter(start, end), {"_id": 0}).sort("created_at", 1).to_list(10000)
    wb = Workbook(); ws = wb.active; ws.title = "Tank Readings"
    headers = ["Date","Time","Session","Tank","pH","pH Target Min","pH Target Max",
               "EC","EC Target Min","EC Target Max","Temp (°C)","Temp Target Min","Temp Target Max",
               "Notes","Recorded By","Created At"]
    ws.append(headers)
    for r in rows:
        ws.append([
            r.get("check_date",""), r.get("check_time",""), r.get("session",""), r.get("tank_id",""),
            r.get("ph_actual"), r.get("ph_target_min"), r.get("ph_target_max"),
            r.get("ec_actual"), r.get("ec_target_min"), r.get("ec_target_max"),
            r.get("temp_actual"), r.get("temp_target_min"), r.get("temp_target_max"),
            r.get("notes",""), r.get("user_name",""), _fmt_dt(r.get("created_at"))
        ])
    _autosize(ws)
    fn = f"Tanks_{start or 'all'}_to_{end or 'all'}.xlsx"
    return _xlsx_response(wb, fn)

@api.get("/export/environment")
async def export_environment(start: Optional[str] = None, end: Optional[str] = None, user=Depends(current_user)):
    rows = await db.environment_readings.find(_date_filter(start, end), {"_id": 0}).sort("created_at", 1).to_list(10000)
    wb = Workbook(); ws = wb.active; ws.title = "Environment"
    ws.append(["Date","Time","Session","Temperature (°C)","Humidity (%)","Notes","Recorded By","Created At"])
    for r in rows:
        ws.append([r.get("check_date",""), r.get("check_time",""), r.get("session",""),
                   r.get("temperature"), r.get("humidity"), r.get("notes",""),
                   r.get("user_name",""), _fmt_dt(r.get("created_at"))])
    # daily averages sheet
    ws2 = wb.create_sheet("Daily Averages")
    ws2.append(["Date","Avg Temperature (°C)","Avg Humidity (%)","Readings"])
    by_day = {}
    for r in rows:
        by_day.setdefault(r.get("check_date",""), []).append(r)
    for d, items in sorted(by_day.items()):
        temps = [x["temperature"] for x in items if x.get("temperature") is not None]
        hums = [x["humidity"] for x in items if x.get("humidity") is not None]
        ws2.append([d,
                    round(sum(temps)/len(temps), 2) if temps else "",
                    round(sum(hums)/len(hums), 2) if hums else "",
                    len(items)])
    _autosize(ws); _autosize(ws2)
    fn = f"Environment_{start or 'all'}_to_{end or 'all'}.xlsx"
    return _xlsx_response(wb, fn)

@api.get("/export/field")
async def export_field(start: Optional[str] = None, end: Optional[str] = None, user=Depends(current_user)):
    rows = await db.field_tasks.find(_date_filter(start, end), {"_id": 0}).sort("created_at", 1).to_list(10000)
    wb = Workbook(); ws = wb.active; ws.title = "Field Tasks"
    ws.append(["Date","Seedling Watered","Seedling pH","Seedling EC",
               "Pest Check","Pest Notes","Leaf Cleaning","Leaf Notes","Notes","Recorded By","Created At"])
    for r in rows:
        ws.append([r.get("check_date",""),
                   "Yes" if r.get("seedling_watered") else "No",
                   r.get("seedling_ph"), r.get("seedling_ec"),
                   "Yes" if r.get("pest_check_done") else "No", r.get("pest_notes",""),
                   "Yes" if r.get("leaf_cleaning_done") else "No", r.get("leaf_notes",""),
                   r.get("notes",""), r.get("user_name",""), _fmt_dt(r.get("created_at"))])
    _autosize(ws)
    fn = f"Field_{start or 'all'}_to_{end or 'all'}.xlsx"
    return _xlsx_response(wb, fn)

@api.get("/export/weekly")
async def export_weekly(start: Optional[str] = None, end: Optional[str] = None, user=Depends(current_user)):
    rows = await db.weekly_checks.find(_date_filter(start, end), {"_id": 0}).sort("created_at", 1).to_list(10000)
    wb = Workbook(); ws = wb.active; ws.title = "Weekly Checks"
    ws.append(["Date","Meter Calibration","Nutrition Qty OK","Nutrition Notes",
               "Tank Filters Cleaned","Notes","Recorded By","Created At"])
    for r in rows:
        ws.append([r.get("check_date",""),
                   "Yes" if r.get("meter_calibration_done") else "No",
                   "Yes" if r.get("nutrition_quantity_ok") else "No", r.get("nutrition_notes",""),
                   "Yes" if r.get("tank_filters_cleaned") else "No",
                   r.get("notes",""), r.get("user_name",""), _fmt_dt(r.get("created_at"))])
    _autosize(ws)
    fn = f"Weekly_{start or 'all'}_to_{end or 'all'}.xlsx"
    return _xlsx_response(wb, fn)

@api.get("/export/monthly")
async def export_monthly(start: Optional[str] = None, end: Optional[str] = None, user=Depends(current_user)):
    rows = await db.monthly_checks.find(_date_filter(start, end), {"_id": 0}).sort("created_at", 1).to_list(10000)
    wb = Workbook(); ws = wb.active; ws.title = "Monthly Checks"
    ws.append(["Date","Tanks Cleaned","Salt Formation OK","Salt Notes",
               "Solution A Qty","Solution B Qty","Solution C Qty","Solutions Ordered",
               "Seeds Qty OK","Seeds Ordered","Notes","Recorded By","Created At"])
    for r in rows:
        ws.append([r.get("check_date",""),
                   "Yes" if r.get("tanks_cleaned") else "No",
                   "Yes" if r.get("salt_formation_ok") else "No", r.get("salt_notes",""),
                   r.get("solution_a_qty"), r.get("solution_b_qty"), r.get("solution_c_qty"),
                   "Yes" if r.get("solutions_ordered") else "No",
                   "Yes" if r.get("seeds_qty_ok") else "No",
                   "Yes" if r.get("seeds_ordered") else "No",
                   r.get("notes",""), r.get("user_name",""), _fmt_dt(r.get("created_at"))])
    _autosize(ws)
    fn = f"Monthly_{start or 'all'}_to_{end or 'all'}.xlsx"
    return _xlsx_response(wb, fn)


@api.get("/")
async def root(): return {"service": "Farmlee Manager API", "status": "ok"}

@api.get("/health")
async def health(): return {"status": "healthy", "time": now_utc().isoformat()}


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
